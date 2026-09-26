# Pruebas del flujo principal del MVP: login con roles, registrar una
# venta con un clic, permisos por rol, y el cálculo de KPIs.

from datetime import date, datetime, time, timedelta, timezone
from io import BytesIO

from openpyxl import load_workbook

from app.kpis.service import ZONA_LIMA, _hoy
from app.ventas.models import Venta
from tests.conftest import crear_producto_prueba, crear_usuario_prueba, login


def test_login_correcto_devuelve_token(client, db_session):
    crear_usuario_prueba(db_session, "duena@negocio.com", "dueno")
    respuesta = client.post(
        "/api/auth/login", data={"username": "duena@negocio.com", "password": "clave123"}
    )
    assert respuesta.status_code == 200
    cuerpo = respuesta.json()
    assert cuerpo["rol"] == "dueno"
    assert cuerpo["access_token"]


def test_login_con_contrasena_incorrecta_falla(client, db_session):
    crear_usuario_prueba(db_session, "duena@negocio.com", "dueno")
    respuesta = client.post(
        "/api/auth/login", data={"username": "duena@negocio.com", "password": "incorrecta"}
    )
    assert respuesta.status_code == 401


def test_vendedor_no_puede_crear_producto(client, db_session):
    crear_usuario_prueba(db_session, "vendedor@negocio.com", "vendedor")
    token = login(client, "vendedor@negocio.com")

    respuesta = client.post(
        "/api/productos",
        json={"nombre": "Ibuprofeno", "precio_venta": 8.0, "stock_actual": 5},
        headers={"Authorization": f"Bearer {token}"},
    )
    assert respuesta.status_code == 403


def test_dueno_puede_crear_producto(client, db_session):
    crear_usuario_prueba(db_session, "duena@negocio.com", "dueno")
    token = login(client, "duena@negocio.com")

    respuesta = client.post(
        "/api/productos",
        json={"nombre": "Ibuprofeno", "precio_venta": 8.0, "stock_actual": 5},
        headers={"Authorization": f"Bearer {token}"},
    )
    assert respuesta.status_code == 201
    assert respuesta.json()["nombre"] == "Ibuprofeno"


def test_registrar_venta_descuenta_stock_y_copia_precio(client, db_session):
    crear_usuario_prueba(db_session, "vendedor@negocio.com", "vendedor")
    producto = crear_producto_prueba(db_session, nombre="Paracetamol", stock=10, precio=5.5)
    token = login(client, "vendedor@negocio.com")

    respuesta = client.post(
        "/api/ventas",
        json={"producto_id": producto.id, "cantidad": 3},
        headers={"Authorization": f"Bearer {token}"},
    )
    assert respuesta.status_code == 201
    venta = respuesta.json()
    assert venta["cantidad"] == 3
    assert venta["precio_unitario"] == 5.5
    assert venta["subtotal"] == 16.5

    # el stock del catálogo debe haber bajado de 10 a 7
    productos = client.get(
        "/api/productos", headers={"Authorization": f"Bearer {token}"}
    ).json()
    assert productos[0]["stock_actual"] == 7


def test_no_se_puede_vender_mas_stock_del_disponible(client, db_session):
    crear_usuario_prueba(db_session, "vendedor@negocio.com", "vendedor")
    producto = crear_producto_prueba(db_session, stock=2)
    token = login(client, "vendedor@negocio.com")

    respuesta = client.post(
        "/api/ventas",
        json={"producto_id": producto.id, "cantidad": 5},
        headers={"Authorization": f"Bearer {token}"},
    )
    assert respuesta.status_code == 400


def test_desactivar_producto_lo_saca_del_catalogo_y_de_venta(client, db_session):
    crear_usuario_prueba(db_session, "vendedor@negocio.com", "vendedor")
    dueno = crear_usuario_prueba(db_session, "duena@negocio.com", "dueno")
    producto = crear_producto_prueba(db_session, nombre="Metamizol", stock=10, precio=3.5)

    token_dueno = login(client, "duena@negocio.com")
    respuesta = client.patch(
        f"/api/productos/{producto.id}/desactivar", headers={"Authorization": f"Bearer {token_dueno}"}
    )
    assert respuesta.status_code == 200
    assert respuesta.json()["activo"] is False

    # catálogo normal (el que usa /venta) ya no lo trae
    catalogo = client.get("/api/productos", headers={"Authorization": f"Bearer {token_dueno}"}).json()
    assert "Metamizol" not in {p["nombre"] for p in catalogo}

    # pero sigue existiendo y se puede ver con incluir_inactivos=true
    catalogo_completo = client.get(
        "/api/productos?incluir_inactivos=true", headers={"Authorization": f"Bearer {token_dueno}"}
    ).json()
    encontrado = next(p for p in catalogo_completo if p["nombre"] == "Metamizol")
    assert encontrado["activo"] is False

    # ya no se puede vender
    token_vendedor = login(client, "vendedor@negocio.com")
    respuesta_venta = client.post(
        "/api/ventas",
        json={"producto_id": producto.id, "cantidad": 1},
        headers={"Authorization": f"Bearer {token_vendedor}"},
    )
    assert respuesta_venta.status_code == 404


def test_reactivar_producto_lo_devuelve_al_catalogo_y_se_puede_vender(client, db_session):
    vendedor = crear_usuario_prueba(db_session, "vendedor@negocio.com", "vendedor")
    dueno = crear_usuario_prueba(db_session, "duena@negocio.com", "dueno")
    producto = crear_producto_prueba(db_session, nombre="Metamizol", stock=10, precio=3.5)
    producto.activo = False
    db_session.commit()

    token_dueno = login(client, "duena@negocio.com")
    respuesta = client.patch(
        f"/api/productos/{producto.id}/reactivar", headers={"Authorization": f"Bearer {token_dueno}"}
    )
    assert respuesta.status_code == 200
    assert respuesta.json()["activo"] is True

    catalogo = client.get("/api/productos", headers={"Authorization": f"Bearer {token_dueno}"}).json()
    assert "Metamizol" in {p["nombre"] for p in catalogo}

    token_vendedor = login(client, "vendedor@negocio.com")
    respuesta_venta = client.post(
        "/api/ventas",
        json={"producto_id": producto.id, "cantidad": 1},
        headers={"Authorization": f"Bearer {token_vendedor}"},
    )
    assert respuesta_venta.status_code == 201


def test_vendedor_no_puede_desactivar_ni_reactivar_producto(client, db_session):
    crear_usuario_prueba(db_session, "vendedor@negocio.com", "vendedor")
    producto = crear_producto_prueba(db_session, nombre="Metamizol", stock=10, precio=3.5)
    token = login(client, "vendedor@negocio.com")

    assert (
        client.patch(f"/api/productos/{producto.id}/desactivar", headers={"Authorization": f"Bearer {token}"}).status_code
        == 403
    )
    assert (
        client.patch(f"/api/productos/{producto.id}/reactivar", headers={"Authorization": f"Bearer {token}"}).status_code
        == 403
    )


def test_desactivar_producto_conserva_su_historial_en_kpis(client, db_session):
    vendedor = crear_usuario_prueba(db_session, "vendedor@negocio.com", "vendedor")
    dueno = crear_usuario_prueba(db_session, "duena@negocio.com", "dueno")
    producto = crear_producto_prueba(db_session, nombre="Metamizol", stock=10, precio=3.5)

    token_vendedor = login(client, "vendedor@negocio.com")
    client.post(
        "/api/ventas",
        json={"producto_id": producto.id, "cantidad": 2},
        headers={"Authorization": f"Bearer {token_vendedor}"},
    )

    token_dueno = login(client, "duena@negocio.com")
    client.patch(f"/api/productos/{producto.id}/desactivar", headers={"Authorization": f"Bearer {token_dueno}"})

    # la venta de hoy (ya con el producto desactivado) debe seguir en los
    # reportes agrupados por producto -- desactivar no debe "romper" nada
    resumen_dia = client.get(
        "/api/kpis/resumen-dia", headers={"Authorization": f"Bearer {token_dueno}"}
    ).json()
    assert any(f["nombre"] == "Metamizol" and f["cantidad_total"] == 2 for f in resumen_dia)

    resumen_mes = client.get(
        "/api/kpis/resumen-mes", headers={"Authorization": f"Bearer {token_dueno}"}
    ).json()
    assert any(f["nombre"] == "Metamizol" and f["unidades_vendidas"] == 2 for f in resumen_mes)

    resumen = client.get("/api/kpis/resumen", headers={"Authorization": f"Bearer {token_dueno}"}).json()
    assert any(p["nombre"] == "Metamizol" for p in resumen["productos_mas_vendidos"])

    # y la línea de unidades por producto (últimos 30 días / jornada de hoy)
    # tampoco debe perder la venta de hoy, aunque el producto ya esté inactivo
    historico_productos = client.get(
        "/api/kpis/historico-productos-diario", headers={"Authorization": f"Bearer {token_dueno}"}
    ).json()
    metamizol_historico = next(p for p in historico_productos if p["nombre"] == "Metamizol")
    assert metamizol_historico["puntos"][-1]["cantidad"] == 2

    jornada = client.get("/api/kpis/jornada-hoy", headers={"Authorization": f"Bearer {token_dueno}"}).json()
    total_unidades_jornada = sum(
        sum(p["unidades_por_periodo"]) + 0 for p in jornada["unidades_por_producto"] if p["nombre"] == "Metamizol"
    )
    assert total_unidades_jornada == 2 or jornada["fuera_de_horario"]["cantidad"] == 2


def test_vendedor_no_puede_ver_historial_completo_de_ventas(client, db_session):
    crear_usuario_prueba(db_session, "vendedor@negocio.com", "vendedor")
    token = login(client, "vendedor@negocio.com")

    respuesta = client.get("/api/ventas", headers={"Authorization": f"Bearer {token}"})
    assert respuesta.status_code == 403

    respuesta_propias = client.get(
        "/api/ventas/mias", headers={"Authorization": f"Bearer {token}"}
    )
    assert respuesta_propias.status_code == 200


def test_kpis_resumen_refleja_las_ventas_del_dia(client, db_session):
    vendedor = crear_usuario_prueba(db_session, "vendedor@negocio.com", "vendedor")
    dueno = crear_usuario_prueba(db_session, "duena@negocio.com", "dueno")
    producto = crear_producto_prueba(db_session, precio=10.0, stock=20)

    token_vendedor = login(client, "vendedor@negocio.com")
    client.post(
        "/api/ventas",
        json={"producto_id": producto.id, "cantidad": 2},
        headers={"Authorization": f"Bearer {token_vendedor}"},
    )

    token_dueno = login(client, "duena@negocio.com")
    respuesta = client.get(
        "/api/kpis/resumen", headers={"Authorization": f"Bearer {token_dueno}"}
    )
    assert respuesta.status_code == 200
    kpis = respuesta.json()
    assert kpis["ventas_hoy_cantidad"] == 2
    assert kpis["ventas_hoy_total"] == 20.0
    assert kpis["productos_mas_vendidos"][0]["nombre"] == producto.nombre


def test_kpis_producto_con_menor_rotacion(client, db_session):
    dueno = crear_usuario_prueba(db_session, "duena@negocio.com", "dueno")
    producto_vendido = crear_producto_prueba(db_session, nombre="Ibuprofeno", stock=50, precio=5.0)
    producto_sin_ventas = crear_producto_prueba(db_session, nombre="Curitas", stock=50, precio=3.0)

    token = login(client, "duena@negocio.com")
    client.post(
        "/api/ventas",
        json={"producto_id": producto_vendido.id, "cantidad": 5},
        headers={"Authorization": f"Bearer {token}"},
    )

    respuesta = client.get("/api/kpis/resumen", headers={"Authorization": f"Bearer {token}"})
    assert respuesta.status_code == 200
    menor_rotacion = respuesta.json()["producto_menor_rotacion"]
    # "Curitas" no tuvo ninguna venta en los últimos 30 días -> debe ganar
    # como el de menor rotación (0 unidades), aunque nunca aparezca en el
    # top de más vendidos.
    assert menor_rotacion["nombre"] == "Curitas"
    assert menor_rotacion["unidades_vendidas"] == 0


def test_kpis_resumen_dia_agrupa_ventas_de_hoy_por_producto(client, db_session):
    vendedor = crear_usuario_prueba(db_session, "vendedor@negocio.com", "vendedor")
    dueno = crear_usuario_prueba(db_session, "duena@negocio.com", "dueno")
    paracetamol = crear_producto_prueba(db_session, nombre="Paracetamol", stock=50, precio=5.0)
    ibuprofeno = crear_producto_prueba(db_session, nombre="Ibuprofeno", stock=50, precio=8.0)

    token_vendedor = login(client, "vendedor@negocio.com")
    # dos ventas de hoy del mismo producto -> deben sumarse en una sola fila
    client.post(
        "/api/ventas",
        json={"producto_id": paracetamol.id, "cantidad": 2},
        headers={"Authorization": f"Bearer {token_vendedor}"},
    )
    client.post(
        "/api/ventas",
        json={"producto_id": paracetamol.id, "cantidad": 3},
        headers={"Authorization": f"Bearer {token_vendedor}"},
    )
    client.post(
        "/api/ventas",
        json={"producto_id": ibuprofeno.id, "cantidad": 1},
        headers={"Authorization": f"Bearer {token_vendedor}"},
    )
    # una venta de ayer no debe aparecer en el resumen de HOY
    db_session.add(
        Venta(
            producto_id=paracetamol.id,
            usuario_id=vendedor.id,
            cantidad=10,
            precio_unitario=5.0,
            subtotal=50.0,
            fecha_venta=datetime.now(timezone.utc) - timedelta(days=1),
        )
    )
    db_session.commit()

    token_dueno = login(client, "duena@negocio.com")
    respuesta = client.get(
        "/api/kpis/resumen-dia", headers={"Authorization": f"Bearer {token_dueno}"}
    )
    assert respuesta.status_code == 200
    filas = {f["nombre"]: f for f in respuesta.json()}

    assert filas["Paracetamol"]["cantidad_total"] == 5
    assert filas["Paracetamol"]["subtotal_total"] == 25.0
    assert filas["Ibuprofeno"]["cantidad_total"] == 1
    assert filas["Ibuprofeno"]["subtotal_total"] == 8.0


def test_kpis_historico_diario_incluye_30_dias_con_ceros(client, db_session):
    dueno = crear_usuario_prueba(db_session, "duena@negocio.com", "dueno")
    producto = crear_producto_prueba(db_session, precio=10.0, stock=20)

    # una venta hace 5 días, ninguna hoy -> el histórico debe mostrar el
    # total en ese día y 0 en todos los demás, sin omitir ninguna fecha
    hace_5_dias = datetime.now(timezone.utc) - timedelta(days=5)
    db_session.add(
        Venta(
            producto_id=producto.id,
            usuario_id=dueno.id,
            cantidad=2,
            precio_unitario=10.0,
            subtotal=20.0,
            fecha_venta=hace_5_dias,
        )
    )
    db_session.commit()

    token = login(client, "duena@negocio.com")
    respuesta = client.get(
        "/api/kpis/historico-diario", headers={"Authorization": f"Bearer {token}"}
    )
    assert respuesta.status_code == 200
    puntos = respuesta.json()

    assert len(puntos) == 30
    totales_por_fecha = {p["fecha"]: p["total"] for p in puntos}
    # la fecha esperada es la de LIMA, no la de UTC (pueden diferir según la
    # hora a la que corra la prueba) -> se usa el mismo criterio que el
    # servicio para no depender de en qué momento del día se ejecuta esto.
    fecha_venta = hace_5_dias.astimezone(ZONA_LIMA).date().isoformat()
    assert totales_por_fecha[fecha_venta] == 20.0
    # todos los demás días deben estar presentes con total 0, no omitidos
    assert sum(1 for t in totales_por_fecha.values() if t == 0.0) == 29


def test_kpis_resumen_mes_agrupa_ventas_del_mes_por_producto(client, db_session):
    vendedor = crear_usuario_prueba(db_session, "vendedor@negocio.com", "vendedor")
    dueno = crear_usuario_prueba(db_session, "duena@negocio.com", "dueno")
    paracetamol = crear_producto_prueba(db_session, nombre="Paracetamol", stock=50, precio=5.0)
    ibuprofeno = crear_producto_prueba(db_session, nombre="Ibuprofeno", stock=50, precio=8.0)

    token_vendedor = login(client, "vendedor@negocio.com")
    client.post(
        "/api/ventas",
        json={"producto_id": paracetamol.id, "cantidad": 2},
        headers={"Authorization": f"Bearer {token_vendedor}"},
    )
    client.post(
        "/api/ventas",
        json={"producto_id": ibuprofeno.id, "cantidad": 1},
        headers={"Authorization": f"Bearer {token_vendedor}"},
    )
    # una venta del mes pasado no debe entrar en el resumen del mes ACTUAL
    db_session.add(
        Venta(
            producto_id=paracetamol.id,
            usuario_id=vendedor.id,
            cantidad=99,
            precio_unitario=5.0,
            subtotal=495.0,
            fecha_venta=datetime.now(timezone.utc) - timedelta(days=45),
        )
    )
    db_session.commit()

    token_dueno = login(client, "duena@negocio.com")
    respuesta = client.get(
        "/api/kpis/resumen-mes", headers={"Authorization": f"Bearer {token_dueno}"}
    )
    assert respuesta.status_code == 200
    filas = {f["nombre"]: f for f in respuesta.json()}

    assert filas["Paracetamol"]["unidades_vendidas"] == 2
    assert filas["Paracetamol"]["total_vendido"] == 10.0
    assert filas["Ibuprofeno"]["unidades_vendidas"] == 1


def test_kpis_historico_productos_diario_desglosa_por_producto_con_ceros(client, db_session):
    dueno = crear_usuario_prueba(db_session, "duena@negocio.com", "dueno")
    vendido = crear_producto_prueba(db_session, nombre="Vendido", stock=50, precio=10.0)
    sin_ventas = crear_producto_prueba(db_session, nombre="SinVentas", stock=50, precio=10.0)

    token = login(client, "duena@negocio.com")
    client.post(
        "/api/ventas",
        json={"producto_id": vendido.id, "cantidad": 3},
        headers={"Authorization": f"Bearer {token}"},
    )

    respuesta = client.get(
        "/api/kpis/historico-productos-diario", headers={"Authorization": f"Bearer {token}"}
    )
    assert respuesta.status_code == 200
    productos = {p["nombre"]: p for p in respuesta.json()}

    assert set(productos.keys()) == {"Vendido", "SinVentas"}
    assert len(productos["Vendido"]["puntos"]) == 30
    assert len(productos["SinVentas"]["puntos"]) == 30
    # el producto sin ventas debe aparecer igual, con todos los puntos en 0
    assert all(p["cantidad"] == 0 for p in productos["SinVentas"]["puntos"])
    # el producto vendido debe tener la venta de hoy reflejada en el último punto
    assert productos["Vendido"]["puntos"][-1]["cantidad"] == 3


def test_kpis_productos_ya_vencidos_es_alerta_separada_de_por_vencer(client, db_session):
    dueno = crear_usuario_prueba(db_session, "duena@negocio.com", "dueno")
    ya_vencido = crear_producto_prueba(db_session, nombre="YaVencido", stock=5, precio=3.0)
    ya_vencido.fecha_vencimiento = date.today() - timedelta(days=10)
    por_vencer = crear_producto_prueba(db_session, nombre="PorVencer", stock=5, precio=3.0)
    por_vencer.fecha_vencimiento = date.today() + timedelta(days=10)
    db_session.commit()

    token = login(client, "duena@negocio.com")
    respuesta = client.get("/api/kpis/resumen", headers={"Authorization": f"Bearer {token}"})
    assert respuesta.status_code == 200
    cuerpo = respuesta.json()

    # "ya vencidos" debe traer al producto vencido, con su stock, y NO al que
    # todavía está por vencer
    nombres_vencidos = {p["nombre"] for p in cuerpo["productos_ya_vencidos"]}
    assert nombres_vencidos == {"YaVencido"}
    assert cuerpo["productos_ya_vencidos"][0]["stock_actual"] == 5

    # "por vencer" (30 días) debe contar solo al que todavía no vence, sin
    # duplicar al que ya está en la alerta de vencidos
    assert cuerpo["productos_por_vencer"] == 1


def test_kpis_dia_usa_hora_local_de_lima_no_utc(client, db_session):
    vendedor = crear_usuario_prueba(db_session, "vendedor@negocio.com", "vendedor")
    dueno = crear_usuario_prueba(db_session, "duena@negocio.com", "dueno")
    venta_tarde = crear_producto_prueba(db_session, nombre="VentaTarde", stock=50, precio=10.0)
    venta_trasnoche = crear_producto_prueba(db_session, nombre="VentaTrasnoche", stock=50, precio=10.0)

    hoy_lima = _hoy()

    # 20:00 hora de Lima de HOY -> en UTC cae ya en la madrugada del día
    # SIGUIENTE (Lima = UTC-5). Con la lógica vieja (basada en UTC) esta
    # venta quedaba excluida de "hoy" hasta el día siguiente -- el bug real
    # que se encontró con datos reales y que corrige este cambio.
    momento_20h_lima_hoy = datetime.combine(hoy_lima, time(20, 0), tzinfo=ZONA_LIMA)
    db_session.add(
        Venta(
            producto_id=venta_tarde.id,
            usuario_id=vendedor.id,
            cantidad=1,
            precio_unitario=10.0,
            subtotal=10.0,
            fecha_venta=momento_20h_lima_hoy.astimezone(timezone.utc),
        )
    )

    # 23:00 hora de Lima de AYER -> en UTC cae ya en la madrugada de HOY.
    # No debe contarse como venta de hoy (era la otra cara del mismo bug).
    momento_23h_lima_ayer = datetime.combine(hoy_lima - timedelta(days=1), time(23, 0), tzinfo=ZONA_LIMA)
    db_session.add(
        Venta(
            producto_id=venta_trasnoche.id,
            usuario_id=vendedor.id,
            cantidad=1,
            precio_unitario=10.0,
            subtotal=10.0,
            fecha_venta=momento_23h_lima_ayer.astimezone(timezone.utc),
        )
    )
    db_session.commit()

    token = login(client, "duena@negocio.com")
    respuesta = client.get("/api/kpis/resumen-dia", headers={"Authorization": f"Bearer {token}"})
    assert respuesta.status_code == 200
    nombres = {f["nombre"] for f in respuesta.json()}

    assert "VentaTarde" in nombres
    assert "VentaTrasnoche" not in nombres


def test_kpis_operaciones_cuenta_transacciones_no_unidades(client, db_session):
    vendedor = crear_usuario_prueba(db_session, "vendedor@negocio.com", "vendedor")
    dueno = crear_usuario_prueba(db_session, "duena@negocio.com", "dueno")
    producto = crear_producto_prueba(db_session, nombre="Paracetamol", stock=50, precio=5.0)

    token = login(client, "vendedor@negocio.com")
    # 3 operaciones separadas -> unidades = 1+2+4 = 7, pero operaciones = 3
    for cantidad in (1, 2, 4):
        client.post(
            "/api/ventas",
            json={"producto_id": producto.id, "cantidad": cantidad},
            headers={"Authorization": f"Bearer {token}"},
        )

    token_dueno = login(client, "duena@negocio.com")

    resumen = client.get("/api/kpis/resumen", headers={"Authorization": f"Bearer {token_dueno}"}).json()
    assert resumen["ventas_hoy_cantidad"] == 7
    assert resumen["ventas_hoy_operaciones"] == 3
    assert resumen["ventas_mes_operaciones"] == 3

    resumen_dia = client.get("/api/kpis/resumen-dia", headers={"Authorization": f"Bearer {token_dueno}"}).json()
    assert resumen_dia[0]["cantidad_total"] == 7
    assert resumen_dia[0]["operaciones"] == 3

    resumen_mes = client.get("/api/kpis/resumen-mes", headers={"Authorization": f"Bearer {token_dueno}"}).json()
    assert resumen_mes[0]["unidades_vendidas"] == 7
    assert resumen_mes[0]["operaciones"] == 3


def test_kpis_jornada_hoy_agrupa_por_periodo_de_3_horas(client, db_session):
    vendedor = crear_usuario_prueba(db_session, "vendedor@negocio.com", "vendedor")
    dueno = crear_usuario_prueba(db_session, "duena@negocio.com", "dueno")
    producto_a = crear_producto_prueba(db_session, nombre="ProductoA", stock=50, precio=10.0)
    producto_b = crear_producto_prueba(db_session, nombre="ProductoB", stock=50, precio=10.0)

    hoy_lima = _hoy()

    def venta_a_hora_lima(producto, hora, cantidad=1):
        momento = datetime.combine(hoy_lima, time(hora, 0), tzinfo=ZONA_LIMA)
        db_session.add(
            Venta(
                producto_id=producto.id,
                usuario_id=vendedor.id,
                cantidad=cantidad,
                precio_unitario=10.0,
                subtotal=10.0 * cantidad,
                fecha_venta=momento.astimezone(timezone.utc),
            )
        )

    venta_a_hora_lima(producto_a, 9)               # periodo 08:00-11:00
    venta_a_hora_lima(producto_b, 13)              # periodo 11:00-14:00
    venta_a_hora_lima(producto_a, 19, cantidad=2)  # periodo 17:00-20:00
    venta_a_hora_lima(producto_a, 6)               # fuera de horario (antes de 08:00)
    venta_a_hora_lima(producto_b, 21)              # fuera de horario (después de 20:00)
    db_session.commit()

    token = login(client, "duena@negocio.com")
    respuesta = client.get("/api/kpis/jornada-hoy", headers={"Authorization": f"Bearer {token}"})
    assert respuesta.status_code == 200
    cuerpo = respuesta.json()

    assert cuerpo["periodos"] == ["08:00–11:00", "11:00–14:00", "14:00–17:00", "17:00–20:00"]
    assert cuerpo["ingresos_por_periodo"] == [10.0, 10.0, 0.0, 20.0]
    assert cuerpo["operaciones_por_periodo"] == [1, 1, 0, 1]

    unidades = {u["nombre"]: u["unidades_por_periodo"] for u in cuerpo["unidades_por_producto"]}
    assert unidades["ProductoA"] == [1, 0, 0, 2]
    assert unidades["ProductoB"] == [0, 1, 0, 0]

    # las ventas fuera de 08:00-20:00 quedan registradas aparte, no
    # descartadas ni mezcladas en los periodos
    assert cuerpo["fuera_de_horario"]["operaciones"] == 2
    assert cuerpo["fuera_de_horario"]["cantidad"] == 2
    assert cuerpo["fuera_de_horario"]["subtotal"] == 20.0


def test_kpis_productos_agotados_no_oculta_a_los_vencidos(client, db_session):
    dueno = crear_usuario_prueba(db_session, "duena@negocio.com", "dueno")
    solo_agotado = crear_producto_prueba(db_session, nombre="SoloAgotado", stock=0, precio=3.0)
    vencido_y_agotado = crear_producto_prueba(db_session, nombre="VencidoYAgotado", stock=0, precio=3.0)
    vencido_y_agotado.fecha_vencimiento = date.today() - timedelta(days=5)
    con_stock = crear_producto_prueba(db_session, nombre="ConStock", stock=10, precio=3.0)
    db_session.commit()

    token = login(client, "duena@negocio.com")
    cuerpo = client.get("/api/kpis/resumen", headers={"Authorization": f"Bearer {token}"}).json()

    nombres_agotados = {p["nombre"] for p in cuerpo["productos_agotados"]}
    nombres_vencidos = {p["nombre"] for p in cuerpo["productos_ya_vencidos"]}

    assert nombres_agotados == {"SoloAgotado", "VencidoYAgotado"}
    assert nombres_vencidos == {"VencidoYAgotado"}
    # el producto vencido-y-agotado debe seguir apareciendo en AMBAS
    # listas: no se ocultan ni se excluyen entre sí
    assert "ConStock" not in nombres_agotados


def test_kpis_exportar_dia_genera_xlsx_solo_de_esa_fecha(client, db_session):
    vendedor = crear_usuario_prueba(db_session, "vendedor@negocio.com", "vendedor")
    dueno = crear_usuario_prueba(db_session, "duena@negocio.com", "dueno")
    producto = crear_producto_prueba(db_session, nombre="Paracetamol", stock=50, precio=5.0)

    hoy_lima = _hoy()
    momento_hoy = datetime.combine(hoy_lima, time(10, 30), tzinfo=ZONA_LIMA)
    momento_ayer = datetime.combine(hoy_lima - timedelta(days=1), time(10, 30), tzinfo=ZONA_LIMA)

    db_session.add_all(
        [
            Venta(
                producto_id=producto.id, usuario_id=vendedor.id, cantidad=3,
                precio_unitario=5.0, subtotal=15.0, fecha_venta=momento_hoy.astimezone(timezone.utc),
            ),
            # una venta de ayer no debe aparecer al exportar HOY
            Venta(
                producto_id=producto.id, usuario_id=vendedor.id, cantidad=99,
                precio_unitario=5.0, subtotal=495.0, fecha_venta=momento_ayer.astimezone(timezone.utc),
            ),
        ]
    )
    db_session.commit()

    token = login(client, "duena@negocio.com")
    respuesta = client.get(
        f"/api/kpis/exportar-dia?fecha={hoy_lima.isoformat()}",
        headers={"Authorization": f"Bearer {token}"},
    )
    assert respuesta.status_code == 200
    assert respuesta.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert f"ventas_{hoy_lima.isoformat()}.xlsx" in respuesta.headers["content-disposition"]

    libro = load_workbook(BytesIO(respuesta.content))
    hoja = libro.active
    filas = list(hoja.iter_rows(values_only=True))

    assert filas[0] == ("Fecha", "Hora", "Producto", "Cantidad", "Subtotal (S/)")
    assert len(filas) == 2  # encabezado + 1 sola venta (la de ayer no debe salir)
    assert filas[1][0] == hoy_lima.isoformat()
    assert filas[1][1] == "10:30:00"
    assert filas[1][2] == "Paracetamol"
    assert filas[1][3] == 3
    assert filas[1][4] == 15.0


def test_no_se_puede_repetir_correo_de_usuario(client, db_session):
    crear_usuario_prueba(db_session, "duena@negocio.com", "dueno")
    token = login(client, "duena@negocio.com")

    respuesta = client.post(
        "/api/usuarios",
        json={
            "nombre_completo": "Otra Persona",
            "correo": "duena@negocio.com",
            "contrasena": "clave123",
            "rol": "vendedor",
        },
        headers={"Authorization": f"Bearer {token}"},
    )
    assert respuesta.status_code == 409
